"""Bounded XLSX writer for server-built saved-recipe reports. Never evaluates formulas."""
import io
import json
import math
import os
import resource
import sys

resource.setrlimit(resource.RLIMIT_CPU, (5, 6))
resource.setrlimit(resource.RLIMIT_AS, (512 * 1024 * 1024, 512 * 1024 * 1024))
os.environ["OPENBLAS_NUM_THREADS"] = "1"
os.environ["OMP_NUM_THREADS"] = "1"

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.worksheet import Worksheet


def text_for_xml(value):
    # Preserve text as text, replacing only characters XML cannot represent.
    return "".join(c if ord(c) in (9, 10, 13) or 0x20 <= ord(c) <= 0xD7FF
                   or 0xE000 <= ord(c) <= 0xFFFD or 0x10000 <= ord(c) <= 0x10FFFF
                   else "\uFFFD" for c in value)


def write_workbook(payload):
    if not isinstance(payload, dict) or set(payload) != {"sheets"}:
        raise ValueError("Invalid report")
    sheets = payload["sheets"]
    if not isinstance(sheets, list) or not 1 <= len(sheets) <= 5:
        raise ValueError("Invalid worksheets")
    workbook = Workbook(write_only=True)
    seen = set()
    cells = 0
    for sheet in sheets:
        if not isinstance(sheet, dict) or set(sheet) != {"name", "rows", "headerRows"}:
            raise ValueError("Invalid worksheet")
        name, rows, headers = sheet["name"], sheet["rows"], sheet["headerRows"]
        if (not isinstance(name, str) or not 1 <= len(name) <= 31
                or any(c in name for c in "\\/*?:[]") or name.casefold() in seen):
            raise ValueError("Invalid worksheet name")
        seen.add(name.casefold())
        if not isinstance(rows, list) or len(rows) > 2000:
            raise ValueError("Too many rows")
        if not isinstance(headers, list) or any(type(h) is not int or h < 1 or h > len(rows) for h in headers):
            raise ValueError("Invalid header rows")
        ws = workbook.create_sheet(name)
        ws.freeze_panes = "A2"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = Worksheet.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        for column in "ABCDEFGHIJKLMNOPQR":
            ws.column_dimensions[column].width = 24
        for row_number, row in enumerate(rows, 1):
            if not isinstance(row, list) or len(row) > 32:
                raise ValueError("Too many columns")
            cells += len(row)
            if cells > 20000:
                raise ValueError("Too many cells")
            output = []
            for value in row:
                if isinstance(value, str):
                    if len(value) > 4000:
                        raise ValueError("Report text too long")
                    cell = WriteOnlyCell(ws, text_for_xml(value))
                    cell.data_type = "s"  # '=...' is literal text, never a formula.
                elif value is None or type(value) in (float, int) and math.isfinite(value):
                    cell = WriteOnlyCell(ws, value)
                    if isinstance(value, float):
                        cell.number_format = "0.000000"
                else:
                    raise ValueError("Invalid report cell")
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if row_number in headers:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="173F35")
                output.append(cell)
            ws.append(output)
    buffer = io.BytesIO()
    workbook.save(buffer)
    result = buffer.getvalue()
    if len(result) > 2 * 1024 * 1024:
        raise ValueError("Report output too large")
    return result


if __name__ == "__main__":
    try:
        raw = sys.stdin.buffer.read(2 * 1024 * 1024 + 1)
        if len(raw) > 2 * 1024 * 1024:
            raise ValueError("Report input too large")
        data = json.loads(raw, parse_constant=lambda _: (_ for _ in ()).throw(ValueError("Nonfinite number")))
        sys.stdout.buffer.write(write_workbook(data))
    except Exception:
        print("Recipe workbook could not be generated.", file=sys.stderr)
        sys.exit(1)
