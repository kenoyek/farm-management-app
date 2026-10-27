"""Bounded XLSX inspection worker. Never evaluates formulas or extracts ZIP files."""
import os
os.environ['OPENBLAS_NUM_THREADS'] = '1'
os.environ['OPENPYXL_DEFUSEDXML'] = 'True'
os.environ['OPENPYXL_LXML'] = 'False'
import resource
resource.setrlimit(resource.RLIMIT_CPU, (6, 7))
resource.setrlimit(resource.RLIMIT_AS, (512 * 1024 * 1024, 512 * 1024 * 1024))
import sys,json,io,zipfile,pathlib,datetime,math,re
from defusedxml.ElementTree import fromstring,iterparse
import openpyxl

class InvalidWorkbook(Exception):
    pass

def inspect(data):
    if len(data)>2*1024*1024:
        raise InvalidWorkbook('Workbook exceeds the 2 MB upload limit.')
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        entries=archive.infolist()
        if len(entries)>2000 or sum(e.file_size for e in entries)>20*1024*1024:
            raise InvalidWorkbook('Workbook archive exceeds inspection limits.')
        names=[e.filename for e in entries]
        if len(names)!=len(set(names)):
            raise InvalidWorkbook('Duplicate workbook archive entries are not supported.')
        if '[Content_Types].xml' not in names or 'xl/workbook.xml' not in names:
            raise InvalidWorkbook('Upload a valid .xlsx workbook.')
        for entry in entries:
            path=pathlib.PurePosixPath(entry.filename)
            if path.is_absolute() or '..' in path.parts or '\\' in entry.filename or entry.flag_bits&1:
                raise InvalidWorkbook('Unsafe or encrypted workbook archive.')
            low=entry.filename.lower()
            if any(part in low for part in ['vbaproject','externallinks/','embeddings/','activex/']):
                raise InvalidWorkbook('Macros, embedded objects and external links are not supported.')
            if low.startswith('xl/worksheets/') and low.endswith('.xml'):
                xml_cells=0
                for _,element in iterparse(archive.open(entry),events=('end',)):
                    tag=element.tag.split('}')[-1]
                    if tag=='row' and int(element.attrib.get('r','0'))>5000:
                        raise InvalidWorkbook('Worksheet row limit exceeded.')
                    if tag=='c':
                        xml_cells+=1
                        ref=element.attrib.get('r','')
                        match=re.fullmatch(r'([A-Z]{1,3})([1-9][0-9]*)',ref)
                        if not match:raise InvalidWorkbook('Invalid workbook cell reference.')
                        column=0
                        for char in match[1]:column=column*26+ord(char)-64
                        if column>64 or int(match[2])>5000 or xml_cells>20000:
                            raise InvalidWorkbook('Worksheet cell limits exceeded.')
                    element.clear()
            if low.endswith('.rels') or low=='[content_types].xml':
                root=fromstring(archive.read(entry))
                if any(node.attrib.get('TargetMode','').lower()=='external' for node in root.iter()):
                    raise InvalidWorkbook('External workbook relationships are not supported.')
                if any('macroenabled' in str(node.attrib).lower() for node in root.iter()):
                    raise InvalidWorkbook('Macro-enabled workbooks are not supported.')
    workbook=openpyxl.load_workbook(io.BytesIO(data),read_only=True,data_only=False,keep_links=False)
    try:
        if len(workbook.worksheets)>12:
            raise InvalidWorkbook('A workbook may contain at most 12 worksheets.')
        sheets=[];area=0;populated=0
        for sheet in workbook:
            if sheet.max_row is None or sheet.max_column is None:sheet.calculate_dimension(force=True)
            if len(sheet.title)>100 or sheet.max_row>5000 or sheet.max_column>64:
                raise InvalidWorkbook('Worksheet limits: 5,000 rows and 64 columns.')
            area+=sheet.max_row*sheet.max_column
            if area>100000:
                raise InvalidWorkbook('Workbook worksheet area exceeds 100,000 cells.')
            cells=[]
            for row in sheet:
                for cell in row:
                    if cell.value is None:continue
                    populated+=1
                    if populated>20000:
                        raise InvalidWorkbook('Workbook exceeds 20,000 populated cells.')
                    value=cell.value
                    if cell.data_type=='f':
                        kind='formula';value=value if isinstance(value,str) else getattr(value,'text','Unsupported formula')
                    elif isinstance(value,bool):kind='boolean'
                    elif isinstance(value,(datetime.datetime,datetime.date,datetime.time)):
                        kind='date';value=value.isoformat()
                    elif isinstance(value,(int,float)):
                        if not math.isfinite(value):raise InvalidWorkbook('Workbook contains a non-finite number.')
                        kind='number'
                    else:kind='text';value=str(value)
                    if len(str(value))>2000 or len(cell.number_format)>2000:
                        raise InvalidWorkbook('Workbook cell text exceeds 2,000 characters.')
                    cells.append({'cell':cell.coordinate,'row':cell.row,'column':cell.column,'kind':kind,'value':value,'format':cell.number_format})
            sheets.append({'name':sheet.title,'rows':sheet.max_row,'columns':sheet.max_column,'cells':cells})
        return {'sheets':sheets}
    finally:workbook.close()

if __name__=='__main__':
    try:
        result=inspect(sys.stdin.buffer.read(2*1024*1024+1))
        output=json.dumps(result,ensure_ascii=True,allow_nan=False)
        if len(output)>4*1024*1024:raise InvalidWorkbook('Workbook inspection output exceeds the limit.')
        print(output)
    except InvalidWorkbook as error:
        print(json.dumps({'error':str(error)}));sys.exit(2)
    except Exception:
        print(json.dumps({'error':'Invalid or unsupported workbook. Save the file as a standard .xlsx workbook.'}));sys.exit(2)
