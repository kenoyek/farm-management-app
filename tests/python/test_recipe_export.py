"""Read generated XLSX files back, checking exact values and literal text cells."""
import io
import json
import subprocess
import sys
import pytest
from openpyxl import load_workbook


def run(payload):
    return subprocess.run([sys.executable, "scripts/export_recipe.py"],
                          input=json.dumps(payload).encode(), capture_output=True, timeout=10)


def test_round_trip_literal_text_numeric_zero_unknown_and_xml_characters():
    result = run({"sheets": [{"name": "Mixing sheet", "headerRows": [1], "rows": [
        ["Ingredient", "kg", "NGN text", "Unknown"],
        ["=HYPERLINK(\"https://invalid.example\")", 75.125, "150.0000", None],
        ["\t@SUM(1)", 0, "0.0000", "Quoted \"name\""],
        ["bad\x01text\ud800", None],
    ]}]})
    assert result.returncode == 0, result.stderr
    wb = load_workbook(io.BytesIO(result.stdout), data_only=False)
    ws = wb["Mixing sheet"]
    assert ws["A2"].data_type == "s"
    assert ws["A2"].value.startswith("=HYPERLINK")
    assert ws["B2"].value == 75.125
    assert ws["C2"].value == "150.0000" and ws["C2"].data_type == "s"
    assert ws["D2"].value is None
    assert ws["B3"].value == 0
    assert ws["A4"].value == "bad\uFFFDtext\uFFFD"
    assert ws.freeze_panes == "A2"
    assert ws["A1"].font.bold
    assert not any(c.data_type == "f" for row in ws for c in row)


@pytest.mark.parametrize("payload", [
    {"sheets": []},
    {"sheets": [{"name": "bad/name", "rows": [], "headerRows": []}]},
    {"sheets": [{"name": "Report", "rows": [[float("nan")]], "headerRows": []}]},
    {"sheets": [{"name": "Report", "rows": [[True]], "headerRows": []}]},
    {"sheets": [{"name": "Report", "rows": [["x" * 4001]], "headerRows": []}]},
    {"sheets": [{"name": "Report", "rows": [[1]] * 2001, "headerRows": []}]},
])
def test_invalid_or_unbounded_input_never_returns_a_workbook(payload):
    result = run(payload)
    assert result.returncode != 0
    assert not result.stdout


def test_full_report_contract_preserves_four_sheets_and_explicit_status():
    sheets = [{"name": name, "headerRows": [1], "rows": [["Status", "DRAFT — NOT APPROVED"], ["Reference UUID", "fixture-id"]]}
              for name in ["Mixing sheet", "Nutrients", "Reference audit", "Specification"]]
    result = run({"sheets": sheets})
    assert result.returncode == 0
    wb = load_workbook(io.BytesIO(result.stdout))
    assert wb.sheetnames == [s["name"] for s in sheets]
    assert wb["Mixing sheet"]["B1"].value == "DRAFT — NOT APPROVED"
